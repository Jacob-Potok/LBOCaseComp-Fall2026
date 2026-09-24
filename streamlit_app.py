from __future__ import annotations

import html
import io
import json
import mimetypes
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

import pandas as pd
import streamlit as st


ROOT = Path(__file__).resolve().parent
APP_FILES = {"streamlit_app.py", "requirements.txt", ".gitignore", "STREAMLIT_SETUP.md"}
IGNORED_DIRS = {".git", ".streamlit", "__pycache__", ".venv", "venv"}
GITHUB_REPO_DEFAULT = "Jacob-Potok/LBOCaseComp-Fall2026"
TEAM_POST_MARKER = "<!-- lbo-case-team-post -->"

st.set_page_config(
    page_title="LBO Case Comp | Team Library",
    page_icon="📚",
    layout="wide",
    initial_sidebar_state="expanded",
)


def get_setting(name: str, default: str = "") -> str:
    """Read configuration from Streamlit Secrets first, then environment."""
    try:
        value = st.secrets.get(name, default)
    except Exception:
        value = default
    return str(value or os.environ.get(name, default))


@st.cache_data(ttl=300, show_spinner=False)
def github_file_update(repo: str, path: str, token: str) -> tuple[str | None, str | None]:
    """Return the latest commit timestamp and URL for a repository file."""
    if not repo:
        return None, None
    url = f"https://api.github.com/repos/{repo}/commits?path={quote(path, safe='')}&per_page=1"
    headers = {
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "lbo-case-comp-library",
    }
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = Request(url, headers=headers)
    try:
        with urlopen(request, timeout=8) as response:
            commits = json.loads(response.read().decode("utf-8"))
        if not commits:
            return None, None
        commit = commits[0]
        timestamp = commit.get("commit", {}).get("committer", {}).get("date")
        return timestamp, commit.get("html_url")
    except (HTTPError, URLError, TimeoutError, ValueError, OSError):
        return None, None


def github_api(repo: str, endpoint: str, token: str, method: str = "GET", payload: dict | None = None):
    """Small GitHub REST helper; the token stays on the Streamlit server."""
    url = f"https://api.github.com/repos/{repo}/{endpoint.lstrip('/')}"
    headers = {
        "Accept": "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
        "User-Agent": "lbo-case-comp-library",
    }
    if token:
        headers["Authorization"] = f"Bearer {token}"
    body = json.dumps(payload).encode("utf-8") if payload is not None else None
    if body is not None:
        headers["Content-Type"] = "application/json"
    request = Request(url, data=body, headers=headers, method=method)
    try:
        with urlopen(request, timeout=12) as response:
            raw = response.read()
            return (json.loads(raw.decode("utf-8")) if raw else {}), None
    except HTTPError as exc:
        detail = ""
        try:
            detail = json.loads(exc.read().decode("utf-8")).get("message", "")
        except Exception:
            pass
        if exc.code == 410:
            return None, "GitHub Issues are disabled for this repository. Enable Issues in repository settings to collect feedback."
        if exc.code in {401, 403, 404}:
            return None, "The app cannot access the feedback store. Check its GitHub token and repository permissions."
        return None, f"GitHub did not accept the request ({exc.code}). {detail}".strip()
    except (URLError, TimeoutError, ValueError, OSError):
        return None, "Could not reach GitHub. Please try again in a moment."


@st.cache_data(ttl=60, show_spinner=False)
def feedback_issues(repo: str, token: str) -> tuple[list[dict], str | None]:
    all_issues: list[dict] = []
    for page in range(1, 11):
        result, error = github_api(
            repo,
            f"issues?state=all&per_page=100&page={page}",
            token,
        )
        if error:
            return [], error
        if not isinstance(result, list):
            return [], "GitHub returned an unexpected response for feedback."
        all_issues.extend(issue for issue in result if "pull_request" not in issue)
        if len(result) < 100:
            break
    return all_issues, None


def feedback_marker(path: str) -> str:
    return f"<!-- lbo-case-feedback:{path} -->"


def issue_for_file(repo: str, path: str, token: str) -> tuple[dict | None, str | None]:
    issues, error = feedback_issues(repo, token)
    if error:
        return None, error
    marker = feedback_marker(path)
    for issue in issues:
        if marker in (issue.get("body") or ""):
            return issue, None
    return None, None


@st.cache_data(ttl=60, show_spinner=False)
def issue_comments(repo: str, issue_number: int, token: str) -> tuple[list[dict], str | None]:
    result, error = github_api(
        repo,
        f"issues/{issue_number}/comments?per_page=100",
        token,
    )
    if error:
        return [], error
    return (result if isinstance(result, list) else []), None


def save_feedback(repo: str, item: dict, token: str, comment_body: str) -> tuple[bool, str]:
    issue, error = issue_for_file(repo, item["relative"], token)
    if error:
        return False, error
    if issue is None:
        issue_title = f"[Site feedback] {item['relative']}"[:240]
        issue_body = (
            f"{feedback_marker(item['relative'])}\n\n"
            f"Feedback thread for **{item['relative']}** on the LBO Case Comp Library website."
        )
        issue, error = github_api(
            repo,
            "issues",
            token,
            method="POST",
            payload={"title": issue_title, "body": issue_body},
        )
        if error:
            return False, error
        feedback_issues.clear()
    _, error = github_api(
        repo,
        f"issues/{issue['number']}/comments",
        token,
        method="POST",
        payload={"body": comment_body},
    )
    if error:
        return False, error
    issue_comments.clear()
    return True, "Thanks — your feedback has been added to this file's discussion."


def render_team_posts(repo: str, token: str) -> None:
    st.markdown("## Team posts")
    st.caption("Share case updates, questions, and ideas with the group. New posts appear first.")

    posts: list[dict] = []
    if token:
        issues, error = feedback_issues(repo, token)
        if error:
            st.warning(error)
        else:
            posts = [
                issue for issue in issues
                if TEAM_POST_MARKER in (issue.get("body") or "")
            ]
            posts.sort(key=lambda issue: issue.get("created_at", ""), reverse=True)
    else:
        st.info("Team posts need the GitHub token in Streamlit app Secrets.")

    if token:
        with st.expander("✍️ Write a team post", expanded=False):
            with st.form("team-post-form", clear_on_submit=True):
                post_title = st.text_input("Title", max_chars=140, placeholder="What should the group know?")
                post_author = st.text_input("Your name (optional)", max_chars=60)
                post_body = st.text_area(
                    "Post",
                    max_chars=8000,
                    height=180,
                    placeholder="Write an update, share an idea, or ask the group a question…",
                    help="You can use simple Markdown for headings, lists, and links.",
                )
                publish = st.form_submit_button("Publish post", type="primary")

            if publish:
                clean_title = post_title.strip()
                clean_body = post_body.strip()
                if not clean_title or not clean_body:
                    st.error("Add both a title and a post before publishing.")
                else:
                    author = html.escape(post_author.strip()) if post_author.strip() else "Guest"
                    issue_body = (
                        f"{TEAM_POST_MARKER}\n\n"
                        f"**Posted by:** {author}\n\n"
                        f"{clean_body}"
                    )
                    with st.spinner("Publishing your post…"):
                        _, error = github_api(
                            repo,
                            "issues",
                            token,
                            method="POST",
                            payload={"title": f"[Team post] {clean_title}"[:240], "body": issue_body},
                        )
                    if error:
                        st.error(error)
                    else:
                        feedback_issues.clear()
                        st.session_state["team_post_notice"] = "Your post is live."
                        st.rerun()

    notice = st.session_state.pop("team_post_notice", None)
    if notice:
        st.success(notice)

    if posts:
        for issue in posts:
            title = issue.get("title", "Team post")
            if title.startswith("[Team post] "):
                title = title.removeprefix("[Team post] ")
            created = issue.get("created_at", "")
            try:
                posted_at = datetime.fromisoformat(created.replace("Z", "+00:00")).astimezone().strftime("%b %-d, %Y · %I:%M %p")
            except ValueError:
                posted_at = ""
            body = (issue.get("body") or "").replace(TEAM_POST_MARKER, "", 1).strip()
            with st.container(border=True):
                st.markdown(f"### {title}")
                if posted_at:
                    st.caption(posted_at)
                st.markdown(body)
    elif token and not error:
        st.info("No team posts yet. Use **Write a team post** to start the conversation.")


def render_feedback(item: dict, repo: str, token: str) -> None:
    st.markdown("#### Comments & suggestions")
    st.caption("Feedback is visible to everyone who can access this site. You can leave a name or post as Guest.")
    notice = st.session_state.get("feedback_notice")
    if notice and notice.get("path") == item["relative"]:
        st.success(notice["message"])
        st.session_state.pop("feedback_notice", None)
    if not token:
        st.info("Comments are not enabled yet. The site owner needs to add a GitHub token with repository Issues read/write access in Streamlit app Secrets.")
        return

    issue, error = issue_for_file(repo, item["relative"], token)
    if error:
        st.warning(error)
        return
    if issue:
        comments, error = issue_comments(repo, issue["number"], token)
        if error:
            st.warning(error)
        elif comments:
            st.caption(f"{len(comments)} comment{'s' if len(comments) != 1 else ''}")
            for comment in comments:
                author = comment.get("user", {}).get("login", "Visitor")
                created = comment.get("created_at", "")
                try:
                    date_label = datetime.fromisoformat(created.replace("Z", "+00:00")).astimezone().strftime("%b %-d, %Y · %I:%M %p")
                except ValueError:
                    date_label = ""
                with st.container(border=True):
                    st.caption(f"{author} · {date_label}" if date_label else author)
                    st.markdown(comment.get("body", ""))
        else:
            st.caption("No feedback yet. Start the discussion for this file.")
    else:
        st.caption("No feedback yet. Start the discussion for this file.")

    with st.form(f"feedback-form-{item['relative']}", clear_on_submit=True):
        st.markdown("**Add feedback**")
        feedback_type = st.radio(
            "Feedback type",
            ["Comment", "Highlight a section + suggest an edit"],
            horizontal=False,
            key=f"feedback-type-{item['relative']}",
        )
        author_name = st.text_input("Your name (optional)", max_chars=60)
        location = st.text_input(
            "Where is it? (optional)",
            placeholder="For example: page 3, slide 5, or LBO Model!C18",
            max_chars=200,
        )
        excerpt = ""
        if feedback_type.startswith("Highlight"):
            excerpt = st.text_area(
                "Paste the passage or section to highlight",
                max_chars=1200,
                height=90,
            )
        message = st.text_area(
            "Your comment or suggested change",
            max_chars=3000,
            height=110,
            placeholder="Share your question, feedback, or proposed revision…",
        )
        submit = st.form_submit_button("Post feedback", type="primary")

    if submit:
        clean_message = message.strip()
        clean_excerpt = excerpt.strip()
        if not clean_message:
            st.error("Add a comment or suggestion before posting.")
            return
        if feedback_type.startswith("Highlight") and not clean_excerpt:
            st.error("Paste the passage you want to highlight so the suggestion has context.")
            return
        last_posted = st.session_state.get("last_feedback_at", 0.0)
        now = datetime.now().timestamp()
        if now - last_posted < 10:
            st.warning("Please wait a few seconds before posting again.")
            return

        kind = "Section suggestion" if feedback_type.startswith("Highlight") else "Comment"
        parts = [f"**{kind}**"]
        display_name = html.escape(author_name.strip()) if author_name.strip() else "Guest"
        parts.append(f"**Submitted by:** {display_name}")
        if location.strip():
            parts.append(f"**Location:** {location.strip()}")
        if clean_excerpt:
            quoted = "\n".join(f"> {line}" for line in clean_excerpt.splitlines())
            parts.append(f"**Highlighted passage:**\n\n{quoted}")
        parts.append(f"**Feedback:**\n\n{clean_message}")
        with st.spinner("Posting your feedback…"):
            success, result_message = save_feedback(repo, item, token, "\n\n".join(parts))
        if success:
            st.session_state["last_feedback_at"] = now
            st.session_state["feedback_notice"] = {
                "path": item["relative"],
                "message": result_message,
            }
            st.rerun()
        else:
            st.error(result_message)


def list_content_files() -> list[dict]:
    items = []
    for path in ROOT.rglob("*"):
        if not path.is_file():
            continue
        relative = path.relative_to(ROOT)
        if any(part in IGNORED_DIRS or part.startswith(".") for part in relative.parts):
            continue
        if path.name in APP_FILES:
            continue
        stat = path.stat()
        items.append(
            {
                "path": path,
                "relative": relative.as_posix(),
                "name": path.name,
                "suffix": path.suffix.lower(),
                "size": stat.st_size,
                "fallback_time": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc),
            }
        )
    return items


def human_size(size: int) -> str:
    value = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.0f} {unit}" if unit == "B" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{value:.1f} GB"


def category_for(suffix: str) -> str:
    groups = {
        "PDF": {".pdf"},
        "Spreadsheets": {".xlsx", ".xls", ".csv", ".tsv"},
        "Presentations": {".pptx", ".ppt"},
        "Documents": {".docx", ".doc", ".md", ".txt", ".rtf"},
        "Images": {".png", ".jpg", ".jpeg", ".gif", ".webp"},
        "Audio / video": {".mp3", ".wav", ".m4a", ".mp4", ".mov", ".webm"},
    }
    for category, suffixes in groups.items():
        if suffix in suffixes:
            return category
    return "Other files"


def parse_timestamp(value: str | None, fallback: datetime) -> datetime:
    if not value:
        return fallback
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return fallback


def read_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def preview_file(item: dict, github_url: str | None) -> None:
    path = item["path"]
    suffix = item["suffix"]
    try:
        data = path.read_bytes()
    except OSError as exc:
        st.error(f"Could not open this file: {exc}")
        return

    mime_type = mimetypes.guess_type(item["name"])[0] or "application/octet-stream"
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if not reader.pages:
                st.info("This PDF has no readable pages.")
            else:
                page_number = st.number_input(
                    "PDF page", min_value=1, max_value=len(reader.pages), value=1,
                    step=1, key=f"page-{item['relative']}",
                )
                text = reader.pages[int(page_number) - 1].extract_text() or ""
                if text.strip():
                    st.text(text)
                else:
                    st.info("This page has no selectable text. Download the PDF to view it.")
                st.caption(f"Page {int(page_number)} of {len(reader.pages)}")
        except Exception as exc:
            st.info(f"PDF text preview is unavailable ({exc}). Download the PDF to view it.")
    elif suffix in {".md", ".markdown"}:
        st.markdown(read_text(data))
    elif suffix in {".txt", ".rtf", ".log"}:
        st.text(read_text(data))
    elif suffix in {".csv", ".tsv"}:
        try:
            separator = "\t" if suffix == ".tsv" else None
            frame = pd.read_csv(io.BytesIO(data), sep=separator, nrows=1000)
            st.dataframe(frame, use_container_width=True, hide_index=True)
            st.caption("Showing up to the first 1,000 rows.")
        except Exception as exc:
            st.info(f"Table preview is unavailable ({exc}). Download the file instead.")
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            tabs = st.tabs(workbook.sheetnames)
            for tab, sheet_name in zip(tabs, workbook.sheetnames):
                with tab:
                    rows = list(workbook[sheet_name].iter_rows(values_only=True, max_row=101))
                    if rows:
                        width = max((len(row) for row in rows), default=0)
                        normalized = [list(row) + [None] * (width - len(row)) for row in rows]
                        frame = pd.DataFrame(normalized[1:], columns=normalized[0])
                        st.dataframe(frame, use_container_width=True, hide_index=True)
                        st.caption("Preview limited to the first 100 rows per sheet.")
                    else:
                        st.info("This sheet is empty.")
            workbook.close()
        except Exception as exc:
            st.info(f"Spreadsheet preview is unavailable ({exc}). Download the workbook instead.")
    elif suffix == ".pptx":
        try:
            from pptx import Presentation

            deck = Presentation(io.BytesIO(data))
            for index, slide in enumerate(deck.slides, start=1):
                lines = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text.strip()
                ]
                with st.expander(f"Slide {index}", expanded=index == 1):
                    st.text("\n\n".join(lines) if lines else "No selectable text on this slide.")
            st.caption(f"{len(deck.slides)} slides")
        except Exception as exc:
            st.info(f"Presentation preview is unavailable ({exc}). Download the PowerPoint instead.")
    elif suffix == ".docx":
        try:
            from docx import Document

            document = Document(io.BytesIO(data))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            st.markdown("\n\n".join(paragraphs) if paragraphs else "_No paragraph text found._")
            for index, table in enumerate(document.tables, start=1):
                st.markdown(f"**Table {index}**")
                st.table([[cell.text for cell in row.cells] for row in table.rows])
        except Exception as exc:
            st.info(f"Word preview is unavailable ({exc}). Download the document instead.")
    elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        st.image(data, use_container_width=True)
    elif suffix in {".mp3", ".wav", ".m4a"}:
        st.audio(data, format=mime_type)
    elif suffix in {".mp4", ".mov", ".webm"}:
        st.video(data, format=mime_type)
    else:
        st.info("This file is available to download. Add a supported file type for an inline preview.")

    download_cols = st.columns([1, 1, 5])
    download_cols[0].download_button(
        "Download file", data=data, file_name=item["name"], mime=mime_type,
        key=f"download-{item['relative']}", use_container_width=True,
    )
    if github_url:
        download_cols[1].link_button("View on GitHub", github_url, use_container_width=True)


st.markdown(
    """
    <style>
      .block-container {max-width: 1180px; padding-top: 2.2rem; padding-bottom: 3rem;}
      .hero {padding: 1.5rem 1.7rem; border-radius: 18px; margin-bottom: 1.2rem;
             background: linear-gradient(120deg, #102a43 0%, #174c5e 58%, #2a746a 100%); color: white;}
      .hero h1 {font-size: 2.15rem; margin: 0 0 .45rem 0; color: white;}
      .hero p {font-size: 1.02rem; margin: 0; color: #e0f2f1;}
      div[data-testid="stExpander"] {border-radius: 12px; border-color: #dbe4ea;}
    </style>
    <div class="hero">
      <h1>📚 LBO Case Comp Library</h1>
      <p>Shared case materials, research, and team updates — newest additions first.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

repo = get_setting("GITHUB_REPO", GITHUB_REPO_DEFAULT).strip()
token = get_setting("GITHUB_TOKEN").strip()
render_team_posts(repo, token)
st.divider()
items = list_content_files()

if repo and items:
    with st.spinner("Checking the latest file updates…"):
        with ThreadPoolExecutor(max_workers=8) as pool:
            results = list(
                pool.map(
                    lambda item: github_file_update(repo, item["relative"], token),
                    items,
                )
            )
    for item, (timestamp, github_url) in zip(items, results):
        item["updated"] = parse_timestamp(timestamp, item["fallback_time"])
        item["github_url"] = github_url
else:
    for item in items:
        item["updated"] = item["fallback_time"]
        item["github_url"] = None

items.sort(key=lambda item: item["updated"], reverse=True)

if not items:
    st.info("No shared files have been added yet. Add a file to the GitHub repository and it will appear here after the app refreshes.")
else:
    categories = ["All files"] + sorted({category_for(item["suffix"]) for item in items})
    control_cols = st.columns([2, 1, 1])
    search = control_cols[0].text_input("Search files", placeholder="Search by file name")
    chosen_category = control_cols[1].selectbox("File type", categories)
    refresh = control_cols[2].button("Refresh feed", use_container_width=True)
    if refresh:
        st.cache_data.clear()
        st.rerun()

    filtered = [
        item for item in items
        if (chosen_category == "All files" or category_for(item["suffix"]) == chosen_category)
        and search.casefold() in item["name"].casefold()
    ]

    stat_cols = st.columns(3)
    stat_cols[0].metric("Files", len(items))
    stat_cols[1].metric("Showing", len(filtered))
    stat_cols[2].metric("Latest update", filtered[0]["updated"].astimezone().strftime("%b %-d, %Y") if filtered else "—")

    if repo and not token:
        st.caption("The list is connected to the GitHub repository. For a private repository, add a GitHub token in Streamlit app Secrets with Contents read-only and Issues read/write access. This enables accurate update dates and persistent comments.")

    if not filtered:
        st.info("No files match that search. Try another name or file type.")
    else:
        for index, item in enumerate(filtered):
            updated = item["updated"].astimezone().strftime("%b %-d, %Y · %I:%M %p")
            label = f"{item['name']}  ·  {category_for(item['suffix'])}  ·  {updated}  ·  {human_size(item['size'])}"
            with st.expander(label, expanded=(index == 0 and not search)):
                if item["relative"] != item["name"]:
                    st.caption(item["relative"])
                preview_file(item, item.get("github_url"))
                st.divider()
                render_feedback(item, repo, token)

st.divider()
st.caption("Files added to the connected GitHub repository appear here after Streamlit redeploys the latest commit. Use the search and file-type filters to find items quickly.")
